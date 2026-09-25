# -*- coding: utf-8 -*-
"""Relatório Fotográfico Portátil - padrão visual baseado no PDF de referência."""
from __future__ import annotations
import io, os, queue, threading
from dataclasses import dataclass
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from PIL import Image, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

APP_TITLE = "Relatório Fotográfico"
BRAND_DARK = "#006B4F"
BRAND_GREEN = "#2F934E"
BRAND_ACCENT = "#F2C500"
BRAND_BG = "#F4F7F5"
BRAND_TEXT = "#173B31"
BRAND_MUTED = "#64756F"
BRAND_BORDER = "#DCE6E1"
UI_FONT = "Segoe UI"
DOC_FONT = "Calibri"
DOC_HEADING_FONT = "Arial"
SUPPORTED_TYPES = (".jpg", ".jpeg", ".png")
A4_WIDTH, A4_HEIGHT = Inches(8.27), Inches(11.69)
MARGIN_X = Inches(0.15)
MARGIN_BOTTOM = Inches(0.50)
HEADER_DISTANCE = Inches(0.18)
FOOTER_DISTANCE = Inches(0.25)
BLUE = "315B78"
GRID = "A7A7A7"

@dataclass
class ReportData:
    date: str
    obra: str
    photos_per_page: int
    image_paths: list[str]

def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=20, start=35, bottom=20, end=35):
    tc_pr = cell._tc.get_or_add_tcPr(); tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}"); tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_table_borders(table, color=GRID, size="7"):
    tbl_pr = table._tbl.tblPr; borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}"); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), size); el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)

def set_cell_width(cell, inches: float):
    tc_pr = cell._tc.get_or_add_tcPr(); tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW"); tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440))); tc_w.set(qn("w:type"), "dxa")

def set_row_height(row, inches: float, exact=True):
    row.height = Inches(inches); row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit"); tr_pr.append(cant)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr(); el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); tr_pr.append(el)

def add_run_font(run, font, size, bold=False, color="222222"):
    run.font.name = font; run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(color)

def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)

def fit_image_to_box(path: str, max_width: float, max_height: float) -> tuple[io.BytesIO, float, float]:
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        w, h = image.size
        ratio = min(max_width * 96 / w, max_height * 96 / h)
        ratio = min(ratio, 1.0)
        nw, nh = max(1, int(w * ratio)), max(1, int(h * ratio))
        if (nw, nh) != image.size:
            image = image.resize((nw, nh), Image.Resampling.LANCZOS)
        out = io.BytesIO(); image.save(out, "JPEG", quality=94, optimize=True); out.seek(0)
        return out, nw / 96, nh / 96

def add_header_table(header, data: ReportData, first_page: bool):
    for p in list(header.paragraphs):
        if not p.text:
            p._element.getparent().remove(p._element)
    table = header.add_table(rows=3 if first_page else 2, cols=1, width=Inches(7.97))
    table.autofit = False; set_table_borders(table, GRID, "8")
    texts = ["RELATÓRIO DE VISITA TÉCNICA", (data.obra or "NOME DA OBRA").upper()]
    if first_page:
        texts.append(f"Registro Fotográfico - {data.date or ''}".strip())
    heights = [0.62, 0.34, 0.30] if first_page else [0.62, 0.34]
    for i, text in enumerate(texts):
        cell = table.cell(i, 0); set_cell_width(cell, 7.97); set_cell_margins(cell, 0, 0, 0, 0); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; clear_paragraph(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        if i == 0: add_run_font(r, DOC_HEADING_FONT, 15, True, BLUE)
        elif i == 1: add_run_font(r, DOC_HEADING_FONT, 12.5, True, BLUE)
        else: add_run_font(r, DOC_FONT, 12, True, "222222")
        table.rows[i].height = Inches(heights[i]); table.rows[i].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    p = header.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)

def configure_section(section, data):
    section.page_width = A4_WIDTH; section.page_height = A4_HEIGHT
    section.left_margin = MARGIN_X; section.right_margin = MARGIN_X; section.bottom_margin = MARGIN_BOTTOM
    section.top_margin = Inches(1.28)
    section.header_distance = HEADER_DISTANCE; section.footer_distance = FOOTER_DISTANCE
    section.different_first_page_header_footer = True
    add_header_table(section.first_page_header, data, True)
    add_header_table(section.header, data, False)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; clear_paragraph(footer)
    r = footer.add_run("Relatório Fotográfico"); add_run_font(r, DOC_FONT, 7, False, "777777")

def add_photo_cell(cell, path, number, image_w, image_h, cell_h):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell, 0, 0, 0, 0)
    p = cell.paragraphs[0]; clear_paragraph(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    stream, w, h = fit_image_to_box(path, image_w, image_h)
    p.add_run().add_picture(stream, width=Inches(w), height=Inches(h))
    cap = cell.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_before = Pt(0); cap.paragraph_format.space_after = Pt(0)
    r = cap.add_run(f"FOTO Nº {number}"); add_run_font(r, DOC_HEADING_FONT, 9.5, True, "222222")

def add_grid(doc, photos, per_page, progress_callback):
    layouts = {2:(1,2),4:(2,2),6:(2,3),8:(2,4)}; cols, rows = layouts[per_page]
    usable_w = 7.97; gap = 0.02; cell_w = (usable_w-gap*(cols-1))/cols
    content_h = 8.80; row_h = (content_h-gap*(rows-1))/rows
    caption_h = 0.28; image_h = max(0.55, row_h-caption_h)
    total=len(photos)
    for start in range(0,total,per_page):
        group=photos[start:start+per_page]
        table=doc.add_table(rows=rows, cols=cols); table.autofit=False; set_table_borders(table, GRID, "8")
        for idx in range(per_page):
            r,c=divmod(idx,cols); row=table.rows[r]; set_row_height(row,row_h,True); cell=row.cells[c]; set_cell_width(cell,cell_w)
            if idx < len(group):
                n=start+idx+1; progress_callback(n,total,Path(group[idx]).name)
                add_photo_cell(cell,group[idx],n,cell_w-0.03,image_h,row_h)
            else:
                set_cell_margins(cell,0,0,0,0)
        if start+per_page<total: doc.add_page_break()

def build_report(data: ReportData, output_path: str, progress_callback):
    doc=Document(); section=doc.sections[0]; configure_section(section,data)
    core=doc.core_properties; core.title="Relatório Fotográfico"; core.subject="Registro Fotográfico"; core.author="Sistema Criador de Relatórios"
    add_grid(doc,data.image_paths,data.photos_per_page,progress_callback)
    doc.save(output_path)

class ReportApp(tk.Tk):
    def __init__(self):
        super().__init__(); self.title(APP_TITLE); self.geometry("1080x720"); self.minsize(920,650); self.configure(bg=BRAND_BG)
        self.photos=[]; self.events=queue.Queue(); self.generating=False; self._configure_styles(); self._build_ui(); self.after(80,self._poll_events)
    def _configure_styles(self):
        s=ttk.Style(self)
        try:s.theme_use("clam")
        except tk.TclError:pass
        s.configure("App.TFrame",background=BRAND_BG)
        s.configure("Card.TFrame",background="#FFFFFF",relief="solid",borderwidth=1)
        s.configure("Title.TLabel",background=BRAND_BG,foreground=BRAND_DARK,font=(UI_FONT,18,"bold"))
        s.configure("Subtitle.TLabel",background=BRAND_BG,foreground=BRAND_MUTED,font=(UI_FONT,9))
        s.configure("CardTitle.TLabel",background="#FFFFFF",foreground=BRAND_DARK,font=(UI_FONT,11,"bold"))
        s.configure("CardText.TLabel",background="#FFFFFF",foreground=BRAND_MUTED,font=(UI_FONT,8))
        s.configure("Field.TLabel",background="#FFFFFF",foreground=BRAND_TEXT,font=(UI_FONT,9,"bold"))
        s.configure("TEntry",padding=(10,8),font=(UI_FONT,10))
        s.configure("TButton",padding=(12,8),font=(UI_FONT,9,"bold"))
        s.configure("Primary.TButton",background=BRAND_DARK,foreground="#FFFFFF",borderwidth=0,padding=(15,10))
        s.map("Primary.TButton",background=[("active",BRAND_GREEN),("disabled","#AAAAAA")],foreground=[("disabled","#EEEEEE")])
        s.configure("Secondary.TButton",background="#FFFFFF",foreground=BRAND_DARK,borderwidth=1,padding=(12,8))
        s.map("Secondary.TButton",background=[("active","#EAF2EE")])
        s.configure("TRadiobutton",background="#FFFFFF",foreground=BRAND_TEXT,font=(UI_FONT,9),padding=4)
        s.configure("Horizontal.TProgressbar",troughcolor="#E5ECE8",background=BRAND_GREEN,bordercolor="#E5ECE8",lightcolor=BRAND_GREEN,darkcolor=BRAND_GREEN,thickness=8)
    def _card(self,parent,title,description,row,column,padx=8,pady=8):
        f=ttk.Frame(parent,style="Card.TFrame",padding=(18,16)); f.grid(row=row,column=column,sticky="nsew",padx=padx,pady=pady)
        ttk.Label(f,text=title,style="CardTitle.TLabel").pack(anchor="w")
        if description:ttk.Label(f,text=description,style="CardText.TLabel").pack(anchor="w",pady=(4,13))
        content=ttk.Frame(f,style="Card.TFrame")
        content.pack(fill="both",expand=True)
        return content
    def _build_ui(self):
        brand=tk.Frame(self,bg=BRAND_DARK,height=68); brand.pack(fill="x"); brand.pack_propagate(False)
        left=tk.Frame(brand,bg=BRAND_DARK); left.pack(side="left",padx=22,pady=10)
        tk.Label(left,text="Departamento de Engenharia",bg=BRAND_DARK,fg="white",font=(UI_FONT,13,"bold")).pack(anchor="w")
        tk.Label(left,text="Sistema Criador de Relatórios",bg=BRAND_DARK,fg="#D8EAE2",font=(UI_FONT,9)).pack(anchor="w",pady=(2,0))
        tk.Frame(self,bg=BRAND_ACCENT,height=3).pack(fill="x")

        outer=ttk.Frame(self,style="App.TFrame",padding=(24,18)); outer.pack(fill="both",expand=True)
        head=ttk.Frame(outer,style="App.TFrame"); head.pack(fill="x",pady=(0,12))
        tk.Label(head,text="▣",bg=BRAND_DARK,fg="white",font=(UI_FONT,14,"bold"),width=3,height=1).pack(side="left",padx=(0,10))
        tb=ttk.Frame(head,style="App.TFrame"); tb.pack(side="left")
        ttk.Label(tb,text="Relatório Fotográfico",style="Title.TLabel").pack(anchor="w")
        ttk.Label(tb,text="Preencha os dados, adicione as fotos e gere seu relatório em Word.",style="Subtitle.TLabel").pack(anchor="w",pady=(2,0))

        body=ttk.Frame(outer,style="App.TFrame"); body.pack(fill="both",expand=True)
        body.columnconfigure(0,weight=1); body.columnconfigure(1,weight=1)
        body.rowconfigure(0,weight=3); body.rowconfigure(1,weight=2)

        data=self._card(body,"Dados do relatório","Informe a data e o nome da obra.",0,0,padx=6,pady=6)
        data.columnconfigure(0,weight=1); data.columnconfigure(1,weight=2)
        self.vars={k:tk.StringVar() for k in ["date","obra"]}
        ttk.Label(data,text="Data",style="Field.TLabel").grid(row=0,column=0,sticky="w",padx=4,pady=(3,4))
        e=ttk.Entry(data,textvariable=self.vars["date"]); e.grid(row=1,column=0,sticky="ew",padx=4,pady=(0,10)); self._add_placeholder(e,"Ex: 25/09/2026")
        ttk.Label(data,text="Nome da Obra",style="Field.TLabel").grid(row=0,column=1,sticky="w",padx=4,pady=(3,4))
        e=ttk.Entry(data,textvariable=self.vars["obra"]); e.grid(row=1,column=1,sticky="ew",padx=4,pady=(0,10)); self._add_placeholder(e,"Ex: Reforma da Escola Municipal")
        tk.Label(data,text="Essas informações aparecerão no cabeçalho do documento.",bg="#FFFFFF",fg=BRAND_MUTED,font=(UI_FONT,8)).grid(row=2,column=0,columnspan=2,sticky="w",padx=4,pady=(8,0))

        photo=self._card(body,"Fotos","Adicione as imagens que farão parte do relatório.",0,1,padx=6,pady=6)
        photo.columnconfigure(0,weight=1); photo.rowconfigure(1,weight=1)
        ttk.Button(photo,text="＋  Adicionar fotos",style="Primary.TButton",command=self.select_photos).grid(row=0,column=0,sticky="w",pady=(0,8))
        self.photo_count_label=ttk.Label(photo,text="0 fotos adicionadas",style="CardText.TLabel"); self.photo_count_label.grid(row=0,column=1,sticky="e",padx=(10,0))
        lf=ttk.Frame(photo,style="Card.TFrame"); lf.grid(row=1,column=0,columnspan=2,sticky="nsew")
        lf.columnconfigure(0,weight=1); lf.rowconfigure(0,weight=1)
        self.photo_list=tk.Listbox(lf,height=7,borderwidth=0,highlightthickness=0,bg="#FAFCFB",fg=BRAND_TEXT,selectbackground="#DDEBE4",selectforeground=BRAND_DARK,activestyle="none",font=(UI_FONT,9))
        self.photo_list.grid(row=0,column=0,sticky="nsew")
        sb=ttk.Scrollbar(lf,orient="vertical",command=self.photo_list.yview); sb.grid(row=0,column=1,sticky="ns")
        self.photo_list.configure(yscrollcommand=sb.set)
        self.photo_empty_label=tk.Label(lf,text="📷\n\nNenhuma foto adicionada\nClique em “＋ Adicionar fotos” para começar.",bg="#FAFCFB",fg=BRAND_MUTED,font=(UI_FONT,9),justify="center")
        self.photo_empty_label.place(relx=0.5,rely=0.5,anchor="center")

        layout=self._card(body,"Layout da página","Escolha quantas fotos serão exibidas em cada página.",1,0,padx=6,pady=6)
        layout.columnconfigure(0,weight=1); layout.columnconfigure(1,weight=1); self.layout_var=tk.IntVar(value=4)
        opts=[(2,"2 fotos","1 × 2"),(4,"4 fotos","2 × 2"),(6,"6 fotos","2 × 3"),(8,"8 fotos","2 × 4")]
        self.layout_boxes={}
        for i,(v,t,sub) in enumerate(opts):
            r,c=divmod(i,2)
            box=tk.Frame(layout,bg="#FFFFFF",highlightbackground=BRAND_BORDER,highlightthickness=1,cursor="hand2")
            box.grid(row=r,column=c,sticky="ew",padx=4,pady=3)
            self.layout_boxes[v]=box
            ttk.Radiobutton(box,text=t,value=v,variable=self.layout_var,command=self._refresh_layout_selection).pack(anchor="w",padx=7,pady=(5,0))
            tk.Label(box,text=sub,bg="#FFFFFF",fg=BRAND_MUTED,font=(UI_FONT,8)).pack(anchor="w",padx=28,pady=(0,5))
            box.bind("<Button-1>",lambda _,value=v:self._choose_layout(value))
        self._refresh_layout_selection()

        action=self._card(body,"Gerar relatório","Quando estiver tudo pronto, crie o documento Word.",1,1,padx=6,pady=6)
        self.progress_label=ttk.Label(action,text="Adicione fotos para começar.",style="CardText.TLabel"); self.progress_label.pack(anchor="w",pady=(2,6))
        self.progress=ttk.Progressbar(action,mode="determinate",maximum=100,style="Horizontal.TProgressbar"); self.progress.pack(fill="x",pady=(0,10))
        self.generate_button=ttk.Button(action,text="▣  Gerar relatório",style="Primary.TButton",command=self.generate_report); self.generate_button.pack(anchor="e")

    def _choose_layout(self,value):
        self.layout_var.set(value)
        self._refresh_layout_selection()

    def _refresh_layout_selection(self):
        for value,box in self.layout_boxes.items():
            selected=value==self.layout_var.get()
            box.configure(highlightbackground=BRAND_DARK if selected else BRAND_BORDER,highlightthickness=2 if selected else 1)

    def _add_placeholder(self,entry,text):
        entry.insert(0,text); entry.configure(foreground=BRAND_MUTED)
        def focus_in(_):
            if entry.get()==text:entry.delete(0,"end");entry.configure(foreground=BRAND_TEXT)
        def focus_out(_):
            if not entry.get():entry.insert(0,text);entry.configure(foreground=BRAND_MUTED)
        entry.bind("<FocusIn>",focus_in); entry.bind("<FocusOut>",focus_out)
    def _real_value(self,key):
        v=self.vars[key].get(); placeholders={"date":"Ex: 25/09/2026","obra":"Ex: Reforma da Escola Municipal"}; return "" if v==placeholders.get(key) else v
    def select_photos(self):
        files=filedialog.askopenfilenames(title="Selecionar fotos",filetypes=[("Imagens", "*.jpg *.jpeg *.png"), ("JPG", "*.jpg *.jpeg"), ("PNG", "*.png")])
        if not files:return
        self.photos=list(files); self.photo_list.delete(0,"end")
        for i,p in enumerate(self.photos,1):self.photo_list.insert("end",f"Foto {i:02d}  •  {Path(p).name}")
        self.photo_empty_label.place_forget()
        self.photo_count_label.config(text=f"{len(self.photos)} fotos adicionadas")
        self.progress_label.config(text=f"{len(self.photos)} fotos prontas para gerar o relatório.")
    def generate_report(self):
        if self.generating:return
        if not self.photos:messagebox.showwarning(APP_TITLE,"Selecione pelo menos uma foto.");return
        out=filedialog.asksaveasfilename(title="Salvar Relatório",defaultextension=".docx",filetypes=[("Documento Word", "*.docx")],initialfile="Relatorio_Fotografico.docx")
        if not out:return
        data=ReportData(self._real_value("date"),self._real_value("obra") or "Nome da obra não informado",self.layout_var.get(),self.photos.copy())
        self.generating=True; self.generate_button.state(["disabled"]); self.progress["value"]=0; self.progress_label.config(text="Iniciando...")
        threading.Thread(target=self._worker,args=(data,out),daemon=True).start()
    def _worker(self,data,out):
        try:
            def cb(n,total,name):self.events.put(("progress",n,total,name))
            build_report(data,out,cb); self.events.put(("done",out))
        except Exception as e:self.events.put(("error",str(e)))
    def _poll_events(self):
        try:
            while True:
                ev=self.events.get_nowait(); kind=ev[0]
                if kind=="progress":
                    _,n,total,name=ev; self.progress["value"]=n/total*100; self.progress_label.config(text=f"Processando foto {n} de {total}...  {name}")
                elif kind=="done":
                    self.generating=False; self.generate_button.state(["!disabled"]); self.progress["value"]=100; self.progress_label.config(text="✓ Relatório concluído com sucesso."); messagebox.showinfo(APP_TITLE,f"Relatório criado com sucesso:\n\n{ev[1]}")
                else:
                    self.generating=False; self.generate_button.state(["!disabled"]); self.progress_label.config(text="Erro ao gerar relatório."); messagebox.showerror(APP_TITLE,f"Não foi possível gerar o relatório:\n\n{ev[1]}")
        except queue.Empty:pass
        self.after(80,self._poll_events)

if __name__ == "__main__":
    ReportApp().mainloop()
